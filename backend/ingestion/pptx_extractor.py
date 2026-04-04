"""
LAKO V2 — PPTX Structure Extractor (Pass 1)

Uses python-pptx for per-slide text extraction.
Each slide is stored as a "page" (slide_num = page_num) in SQLite.
Image-heavy slides are flagged in the log for the future vision pipeline.
No LLM calls — pure deterministic extraction.
"""

import logging
from pathlib import Path
from typing import Callable, Optional

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from db.document_store import insert_page

logger = logging.getLogger(__name__)

# Shape types that indicate a visual/diagram is present on the slide
_VISUAL_SHAPE_TYPES = {
    MSO_SHAPE_TYPE.PICTURE,
    MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT,
    MSO_SHAPE_TYPE.LINKED_OLE_OBJECT,
    MSO_SHAPE_TYPE.CHART,
    MSO_SHAPE_TYPE.GROUP,
}

# Placeholder index 0 is always the slide title
_TITLE_PLACEHOLDER_IDX = 0


def _slide_has_visuals(slide) -> bool:
    """Return True if the slide contains any picture, chart, or diagram shape."""
    for shape in slide.shapes:
        try:
            if shape.shape_type in _VISUAL_SHAPE_TYPES:
                return True
        except Exception:
            pass
    return False


def _extract_slide_content(slide) -> tuple[str, list[str]]:
    """
    Extract all text from a slide's shapes and notes.

    Returns:
        raw_text: Full concatenated slide text (shapes + notes).
        headers:  List of title-placeholder strings (at most one per slide).
    """
    text_parts: list[str] = []
    headers:    list[str] = []

    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue

        shape_text = shape.text_frame.text.strip()
        if not shape_text:
            continue

        # Detect title placeholder
        if shape.is_placeholder:
            try:
                if shape.placeholder_format.idx == _TITLE_PLACEHOLDER_IDX:
                    headers.append(shape_text)
            except Exception:
                pass

        text_parts.append(shape_text)

    # Speaker notes
    try:
        if slide.has_notes_slide:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()
            if notes_text:
                text_parts.append(f"[Notes: {notes_text}]")
    except Exception:
        pass

    return "\n".join(text_parts), headers


async def extract_pptx(
    file_path: str,
    doc_id: str,
    progress_callback: Optional[Callable[[int, int], None]] = None,
) -> int:
    """
    Extract all slides from *file_path* into the SQLite document store.

    Args:
        file_path: Absolute path to the PPTX file.
        doc_id:    UUID assigned to this document in the documents table.
        progress_callback: Optional fn(slides_done, total_slides).

    Returns:
        Total number of slides extracted.
    """
    path = Path(file_path)

    try:
        prs = Presentation(str(path))
    except Exception as exc:
        logger.error("Cannot open PPTX %s: %s", path.name, exc)
        raise

    total_slides = len(prs.slides)
    logger.info("Extracting PPTX %s — %d slides", path.name, total_slides)

    slides_done = 0
    for idx, slide in enumerate(prs.slides):
        slide_num = idx + 1      # 1-indexed, matches page_num convention

        raw_text, headers = _extract_slide_content(slide)
        char_count        = len(raw_text)
        has_visuals       = _slide_has_visuals(slide)

        if has_visuals and char_count < 50:
            logger.debug(
                "Slide %d of %s is image-heavy (< 50 chars text) — flagged for vision pipeline",
                slide_num, path.name,
            )

        await insert_page(
            doc_id=doc_id,
            page_num=slide_num,
            raw_text=raw_text,
            headers=headers,
            has_tables=False,    # PPTX table detection deferred to retrieval
            char_count=char_count,
        )

        slides_done += 1
        if progress_callback:
            progress_callback(slides_done, total_slides)

    logger.info("PPTX extraction complete: %s — %d slides stored", path.name, slides_done)
    return slides_done
