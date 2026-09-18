"""
LAKO — XLSX Pipeline (V3 Direct-RAG)

Processes Excel workbooks using openpyxl.
Each worksheet is treated as a "page".

Strategy per sheet:
  - If the sheet has headers in row 1: flatten to "Col: Val | Col: Val" prose rows.
  - If no clear headers: concatenate cell values row by row.
  - Each row is kept as a single logical unit (never split mid-row).
  - smart_chunk() is then applied to the flattened text for long sheets.

No LLM calls — pure deterministic extraction.
"""

import logging
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable, Optional

import openpyxl
import structlog

logger = structlog.get_logger(__name__)

ProgressCallback = Optional[Callable[[str, str, int, int], None]]


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _cell_value(cell) -> str:
    """Return a clean string for any cell value."""
    if cell.value is None:
        return ""
    return str(cell.value).strip()


def _flatten_sheet(ws) -> str:
    """
    Flatten a worksheet to prose rows.

    Row 1 is assumed to be the header row.
    Each subsequent row becomes:
      "Header1: Val1 | Header2: Val2 | ..."
    Empty rows are skipped.
    """
    rows = list(ws.iter_rows())
    if not rows:
        return ""

    headers = [_cell_value(cell) for cell in rows[0]]
    has_headers = any(h for h in headers)

    prose_rows: list[str] = []

    for row in rows[1:] if has_headers else rows:
        cells = [_cell_value(c) for c in row]
        if not any(cells):
            continue   # skip blank rows

        if has_headers:
            pairs = [
                f"{h}: {v}" if h else v
                for h, v in zip(headers, cells)
                if v
            ]
        else:
            pairs = [v for v in cells if v]

        if pairs:
            prose_rows.append(" | ".join(pairs))

    return "\n".join(prose_rows)


async def xlsx_pipeline(
    file_path: str,
    doc_id: str,
    filename: str,
    progress_callback: ProgressCallback = None,
) -> list[tuple[str, dict]]:
    """
    Process an XLSX file into (text, metadata) chunks.

    Each worksheet becomes one or more chunks.
    Chunk metadata page number = worksheet position (1-indexed).
    """
    from ingestion.service import smart_chunk

    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    except Exception as exc:
        logger.error("xlsx_open_failed", filename=filename, error=str(exc))
        raise

    sheet_names = wb.sheetnames
    total_sheets = len(sheet_names)
    chunks: list[tuple[str, dict]] = []

    logger.info(
        "xlsx_pipeline_start",
        filename=filename, sheets=total_sheets, doc_id=doc_id,
    )

    for sheet_idx, sheet_name in enumerate(sheet_names):
        sheet_num = sheet_idx + 1
        ws        = wb[sheet_name]

        meta = {
            "doc_id":      doc_id,
            "filename":    filename,
            "page":        sheet_num,
            "sheet_name":  sheet_name,
            "bbox":        [0, 0, 0, 0],
            "font_size":   0.0,
            "is_bold":     False,
            "block_type":  "table",
            "ingested_at": _now(),
            "source_type": "xlsx",
        }

        try:
            text = _flatten_sheet(ws)
            if text.strip():
                sheet_chunks = smart_chunk(text, sheet_num, meta, max_tokens=400, overlap=50)
                chunks.extend(sheet_chunks)
                logger.debug(
                    "xlsx_sheet_processed",
                    sheet=sheet_name, chunks=len(sheet_chunks), doc_id=doc_id,
                )
            else:
                logger.debug("xlsx_sheet_empty", sheet=sheet_name, doc_id=doc_id)

        except Exception as exc:
            logger.error(
                "xlsx_sheet_failed",
                sheet=sheet_name, doc_id=doc_id, error=str(exc),
            )
            # Continue to next sheet

        if progress_callback:
            progress_callback(doc_id, "parsing", sheet_num, total_sheets)

    wb.close()

    logger.info(
        "xlsx_pipeline_complete",
        filename=filename, doc_id=doc_id, sheets=total_sheets, chunks=len(chunks),
    )
    return chunks
